"""Integration tests for ``app.scripts.reingest_bim_model``.

These tests exercise the BIM re-ingest CLI end-to-end against an
isolated temp SQLite (set up by the shared ``tests/conftest.py``).  DDC
is mocked by default — we stub :func:`process_ifc_file` to feed
pre-canned element dicts in, so the suite stays under one second.

A single opt-in test (``@pytest.mark.slow``) runs the *real* DDC
RvtExporter against the showcase ``c5436288`` model when both the
binary and the blob are present in the local environment.  Skipped
silently when either is missing so CI without DDC stays green.
"""

from __future__ import annotations

import asyncio
import json
import os
import shutil
import sys
import tempfile
import uuid
from datetime import UTC, datetime
from pathlib import Path
from typing import Any
from unittest.mock import patch

import pytest


# ── Real DB setup ────────────────────────────────────────────────────────────
# We need actual SQLite tables for these tests; conftest pre-imports the
# models so create_all sees them.

@pytest.fixture(scope="module", autouse=True)
def _create_tables() -> None:
    """Create all ORM tables in the per-session SQLite."""
    from sqlalchemy import create_engine

    from app.config import get_settings
    from app.database import Base

    sync_url = get_settings().database_sync_url
    engine = create_engine(sync_url)
    Base.metadata.create_all(engine)
    engine.dispose()


# ── Helpers ──────────────────────────────────────────────────────────────────


def _make_placeholder_xlsx(path: Path, n_rows: int = 380) -> None:
    """Write a synthetic 'snapshot-style' xlsx with placeholder rows.

    Shape matches the v3.x bulk-loaded RVT placeholders (no ID col, no
    UniqueId, etc.) — these will produce mesh_ref=None and trivial
    names like 'Walls 1'.  Used to seed DB rows for the 'before' state.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = ["Category : String", "Name : String"]
    ws.append(headers)
    for i in range(1, n_rows + 1):
        ws.append(["OST_Walls", f"Walls {i}"])
    wb.save(path)
    wb.close()


def _make_real_xlsx(path: Path, n_rows: int = 1500) -> None:
    """Write a synthetic 'DDC v18 output' xlsx with real columns.

    Each row has the populated DDC columns (ID, UniqueId, Type Name,
    Category, Family, Level, …) so the resulting BIMElements will have
    populated ``mesh_ref`` and meaningful ``properties``.  Acts as a
    stand-in for the actual DDC output to keep the test fast.
    """
    import openpyxl

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Sheet1"
    headers = [
        "ID", "UniqueId : String", "Category : String",
        "Family : String", "Family Name : String", "Type Name : String",
        "Name : String", "Level : String", "Mark : String",
        "Width : Double", "Height : Double", "Volume : Double",
    ]
    ws.append(headers)
    families = ["Basic Wall", "Curtain Wall", "M_Single-Flush", "Basic Floor", "Basic Roof"]
    type_names = [
        "Exterior - Brick on CMU", "Storefront", "0915 x 2134mm",
        "Generic 250mm", "Generic - 400mm",
    ]
    categories = ["OST_Walls", "OST_Doors", "OST_Floors", "OST_Roofs", "OST_Walls"]
    levels = ["Level 1", "Level 2", "Roof"]
    for i in range(n_rows):
        elem_id = 100000 + i
        fam = families[i % len(families)]
        tn = type_names[i % len(type_names)]
        cat = categories[i % len(categories)]
        # Synthetic Revit UniqueId: <GUID>-<hex element id>
        uniq = f"deadbeef-cafe-1234-5678-90abcdef0000-{elem_id:08x}"
        ws.append([
            elem_id, uniq, cat, fam, fam, tn, tn,
            levels[i % len(levels)], f"M-{i:04d}",
            1200.0 + i, 2400.0, (1.2 + i * 0.01),
        ])
    wb.save(path)
    wb.close()


def _seed_placeholder_model(
    project_id: uuid.UUID, n_rows: int = 380,
) -> tuple[uuid.UUID, Path]:
    """Insert a BIMModel + ``n_rows`` placeholder BIMElement rows.

    Also drops a fake ``original.rvt`` blob into the storage backend so
    the script's ``_resolve_original_path`` succeeds (the file body
    doesn't matter — DDC is mocked).

    Returns ``(model_id, original_path_on_disk)``.
    """
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session

    from app.config import get_settings
    from app.modules.bim_hub import file_storage as bim_file_storage
    from app.modules.bim_hub.models import BIMElement, BIMModel

    sync_url = get_settings().database_sync_url
    engine = create_engine(sync_url)
    model_id = uuid.uuid4()
    with Session(engine) as session:
        model = BIMModel(
            id=model_id,
            project_id=project_id,
            name=f"Synth model {model_id}",
            model_format="rvt",
            status="ready",
            element_count=n_rows,
            storey_count=6,
            canonical_file_path=(
                f"bim/{project_id}/{model_id}/geometry.glb"
            ),
        )
        session.add(model)
        for i in range(n_rows):
            session.add(
                BIMElement(
                    model_id=model_id,
                    stable_id=f"RVT-synth-{i:04d}",
                    element_type="Walls",
                    name=f"Walls {i + 1}",
                    storey="Level 1",
                    discipline="architecture",
                    properties={},
                    quantities={},
                    mesh_ref=None,  # the bug
                )
            )
        session.commit()
    engine.dispose()

    # Drop the original.rvt blob.  LocalStorageBackend writes synchronously.
    backend = bim_file_storage._backend()  # noqa: SLF001 — test helper
    key = bim_file_storage.original_cad_key(project_id, model_id, "rvt")
    path = backend._path_for(key)  # noqa: SLF001
    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_bytes(b"FAKE-RVT-PAYLOAD")
    return model_id, path


def _count_rows_with_mesh_ref(model_id: uuid.UUID) -> tuple[int, int]:
    """Return ``(total_rows, rows_with_non_null_mesh_ref)``."""
    from sqlalchemy import create_engine, func, select

    from app.config import get_settings
    from app.modules.bim_hub.models import BIMElement

    sync_url = get_settings().database_sync_url
    engine = create_engine(sync_url)
    with engine.connect() as conn:
        total = conn.execute(
            select(func.count(BIMElement.id)).where(BIMElement.model_id == model_id)
        ).scalar()
        with_mesh = conn.execute(
            select(func.count(BIMElement.id))
            .where(BIMElement.model_id == model_id)
            .where(BIMElement.mesh_ref.isnot(None))
        ).scalar()
    engine.dispose()
    return int(total or 0), int(with_mesh or 0)


def _fetch_sample_rows(model_id: uuid.UUID, limit: int = 50) -> list[dict[str, Any]]:
    """Fetch a few rows for sample-based assertions."""
    from sqlalchemy import create_engine, select

    from app.config import get_settings
    from app.modules.bim_hub.models import BIMElement

    sync_url = get_settings().database_sync_url
    engine = create_engine(sync_url)
    with engine.connect() as conn:
        rows = conn.execute(
            select(
                BIMElement.stable_id, BIMElement.name,
                BIMElement.mesh_ref, BIMElement.properties,
            )
            .where(BIMElement.model_id == model_id)
            .limit(limit)
        ).all()
    engine.dispose()
    return [
        {
            "stable_id": r[0], "name": r[1],
            "mesh_ref": r[2], "properties": r[3] or {},
        }
        for r in rows
    ]


# ── Mock factory ─────────────────────────────────────────────────────────────


def _make_mock_process_ifc_file(xlsx_template: Path, dae_text: str | None = None):
    """Return a stub for :func:`process_ifc_file` that runs the real
    column-mapping pipeline against ``xlsx_template``.

    By calling through :func:`parse_cad_excel` +
    :func:`_excel_elements_to_bim_result` we exercise the production
    mapping code; only the DDC subprocess invocation is bypassed.
    """
    def _stub(ifc_path: Path, output_dir: Path, depth: str = "standard") -> dict[str, Any]:
        from app.modules.bim_hub.ifc_processor import _excel_elements_to_bim_result
        from app.modules.boq.cad_import import parse_cad_excel

        raw_elements = parse_cad_excel(xlsx_template)

        real_dae_path: Path | None = None
        if dae_text:
            real_dae_path = output_dir / "geometry.dae"
            real_dae_path.write_text(dae_text, encoding="utf-8")

        return _excel_elements_to_bim_result(
            raw_elements,
            output_dir,
            real_dae_path=real_dae_path,
        )
    return _stub


# ── Tests ────────────────────────────────────────────────────────────────────


@pytest.fixture
def synthetic_xlsx(tmp_path: Path) -> Path:
    """Pre-built 'DDC v18 output' xlsx with real-shaped columns."""
    xlsx = tmp_path / "ddc_real.xlsx"
    _make_real_xlsx(xlsx, n_rows=1500)
    return xlsx


def _run_cli(*args: str) -> int:
    """Invoke ``app.scripts.reingest_bim_model.main`` directly."""
    from app.scripts.reingest_bim_model import main
    return main(list(args))


def test_reingest_validates_model_exists(tmp_path: Path) -> None:
    """A non-existent UUID should exit non-zero with a clear error."""
    bogus = str(uuid.uuid4())
    rc = _run_cli(bogus)
    assert rc == 1


def test_reingest_missing_rvt_blob() -> None:
    """When the DB row exists but the original blob is gone, exit 2."""
    from sqlalchemy import create_engine
    from sqlalchemy.orm import Session

    from app.config import get_settings
    from app.modules.bim_hub.models import BIMModel

    sync_url = get_settings().database_sync_url
    engine = create_engine(sync_url)
    model_id = uuid.uuid4()
    project_id = uuid.uuid4()
    with Session(engine) as session:
        session.add(
            BIMModel(
                id=model_id,
                project_id=project_id,
                name="No-blob model",
                model_format="rvt",
                status="ready",
                element_count=0,
                storey_count=0,
            )
        )
        session.commit()
    engine.dispose()
    # Deliberately do NOT seed the storage blob.
    rc = _run_cli(str(model_id))
    assert rc == 2


def test_reingest_replaces_synthetic_rows_with_real_data(
    synthetic_xlsx: Path,
) -> None:
    """Happy path: 380 placeholder rows → real-data rows with mesh_ref."""
    project_id = uuid.uuid4()
    model_id, _orig = _seed_placeholder_model(project_id, n_rows=380)

    # Sanity: pre-state matches the bug description.
    before_total, before_with_mesh = _count_rows_with_mesh_ref(model_id)
    assert before_total == 380
    assert before_with_mesh == 0  # all NULL mesh_ref

    # Patch the DDC pass with our xlsx-driven stub.
    stub = _make_mock_process_ifc_file(synthetic_xlsx)
    with patch(
        "app.modules.bim_hub.ifc_processor.process_ifc_file", side_effect=stub
    ):
        rc = _run_cli(str(model_id))
    assert rc == 0

    after_total, after_with_mesh = _count_rows_with_mesh_ref(model_id)
    # New row count > 1000 (the bug remediation threshold from instructions)
    assert after_total > 1000, f"only {after_total} rows after reingest"
    # 100% mesh_ref populated
    assert after_with_mesh == after_total, (
        f"{after_total - after_with_mesh} rows still missing mesh_ref"
    )

    # Sample names should no longer match the "Walls N" placeholder
    # pattern, and at least 50% must carry family + type_name in props.
    sample = _fetch_sample_rows(model_id, limit=200)
    placeholder_re = {f"Walls {i}" for i in range(1, 65)}
    placeholder_hits = sum(1 for r in sample if r["name"] in placeholder_re)
    assert placeholder_hits == 0, (
        f"{placeholder_hits}/{len(sample)} rows still have the 'Walls N' "
        f"placeholder name"
    )
    with_meta = sum(
        1 for r in sample
        if (r["properties"].get("family") and r["properties"].get("type_name"))
    )
    assert with_meta / len(sample) >= 0.5, (
        f"only {with_meta}/{len(sample)} rows have family+type_name"
    )


def test_reingest_idempotent(synthetic_xlsx: Path) -> None:
    """Running re-ingest twice must yield identical row counts (no dupes)."""
    project_id = uuid.uuid4()
    model_id, _ = _seed_placeholder_model(project_id, n_rows=380)
    stub = _make_mock_process_ifc_file(synthetic_xlsx)

    with patch(
        "app.modules.bim_hub.ifc_processor.process_ifc_file", side_effect=stub
    ):
        rc1 = _run_cli(str(model_id))
        assert rc1 == 0
        n1, _ = _count_rows_with_mesh_ref(model_id)

        rc2 = _run_cli(str(model_id))
        assert rc2 == 0
        n2, _ = _count_rows_with_mesh_ref(model_id)

    assert n1 == n2, f"row count changed across reingest calls: {n1} → {n2}"


def test_reingest_dry_run_no_mutation(synthetic_xlsx: Path) -> None:
    """--dry-run must leave the DB untouched."""
    project_id = uuid.uuid4()
    model_id, _ = _seed_placeholder_model(project_id, n_rows=380)
    before_total, before_with_mesh = _count_rows_with_mesh_ref(model_id)
    assert before_total == 380

    stub = _make_mock_process_ifc_file(synthetic_xlsx)
    with patch(
        "app.modules.bim_hub.ifc_processor.process_ifc_file", side_effect=stub
    ):
        rc = _run_cli(str(model_id), "--dry-run")
    assert rc == 0

    after_total, after_with_mesh = _count_rows_with_mesh_ref(model_id)
    assert after_total == before_total
    assert after_with_mesh == before_with_mesh


def test_reingest_backup_creates_file(synthetic_xlsx: Path) -> None:
    """--backup-rows must dump the OLD rows to data/reingest_backup_*.json."""
    project_id = uuid.uuid4()
    model_id, _ = _seed_placeholder_model(project_id, n_rows=380)

    # The script writes under <repo>/data/.  Snapshot any existing files
    # so we can find the one this run produced.  parents[3] mirrors what
    # the script's own resolution does (scripts/ → app/ → backend/ → repo).
    data_dir = Path(__file__).resolve().parents[3] / "data"
    before = set(data_dir.glob(f"reingest_backup_{model_id}_*.json"))

    stub = _make_mock_process_ifc_file(synthetic_xlsx)
    with patch(
        "app.modules.bim_hub.ifc_processor.process_ifc_file", side_effect=stub
    ):
        rc = _run_cli(str(model_id), "--backup-rows")
    assert rc == 0

    after = set(data_dir.glob(f"reingest_backup_{model_id}_*.json"))
    new_files = after - before
    assert len(new_files) == 1, (
        f"expected one new backup, got {len(new_files)}: {new_files}"
    )
    backup_path = next(iter(new_files))
    try:
        payload = json.loads(backup_path.read_text(encoding="utf-8"))
        assert isinstance(payload, list)
        # Backed-up rows = the 380 placeholder rows we seeded.
        assert len(payload) == 380
        # All backed-up rows had mesh_ref=None (the bug state)
        assert all(row["mesh_ref"] is None for row in payload)
        # Sanity: each row carries stable_id + name like 'Walls N'.
        assert payload[0]["stable_id"].startswith("RVT-synth-")
        assert payload[0]["name"] == "Walls 1"
    finally:
        try:
            backup_path.unlink()
        except OSError:
            pass


# ── Real-DDC smoke test ──────────────────────────────────────────────────────


_SHOWCASE_MODEL_ID = "c5436288-8f71-5d89-95a4-c2a4372a5cb3"
_SHOWCASE_PROJECT_ID = "24c3ddfb-00db-44f0-b0b8-9d7ad4078cf2"
_SHOWCASE_RVT = (
    Path(__file__).resolve().parents[3]
    / "data" / "bim" / _SHOWCASE_PROJECT_ID / _SHOWCASE_MODEL_ID / "original.rvt"
)


@pytest.mark.slow
@pytest.mark.skipif(
    not _SHOWCASE_RVT.exists(),
    reason="Showcase RVT blob not present in this checkout",
)
def test_reingest_one_real_showcase_model() -> None:
    """Run the real DDC RvtExporter against the c5436288 showcase model.

    This is the production scenario the script was written for; it's
    marked ``@slow`` because each DDC pass takes ~60-90 s.  When the
    test runs it (a) creates a fresh DB row + copies the blob to the
    test storage backend, (b) invokes the CLI for real, (c) asserts
    that the bug pattern (NULL mesh_ref, 'Walls N' names) is gone.
    """
    from app.modules.boq.cad_import import find_converter

    if find_converter("rvt") is None:
        pytest.skip("DDC RvtExporter not installed in this environment")

    project_id = uuid.UUID(_SHOWCASE_PROJECT_ID)
    model_id, target_blob = _seed_placeholder_model(project_id, n_rows=380)
    # Replace the fake blob with the real one.
    shutil.copyfile(_SHOWCASE_RVT, target_blob)

    rc = _run_cli(str(model_id))
    assert rc == 0

    total, with_mesh = _count_rows_with_mesh_ref(model_id)
    assert total > 1000
    assert with_mesh == total
    sample = _fetch_sample_rows(model_id, limit=200)
    placeholder_hits = sum(
        1 for r in sample if r["name"] and r["name"].startswith("Walls ")
        and r["name"][6:].isdigit()
    )
    assert placeholder_hits == 0
